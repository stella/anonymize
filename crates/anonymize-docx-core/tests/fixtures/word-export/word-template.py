# Builds word-template.docx from python-docx's default template, a package
# saved by Microsoft Word, adding synthetic text in Word's built-in styles.
# Usage: python3 word-template.py <output.docx> (python-docx 1.2).
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"

document = Document()
properties = document.core_properties
properties.author = "Synthetic Metadata Author"
properties.last_modified_by = "Synthetic Metadata Author"
properties.title = "Synthetic Metadata Title"
properties.subject = "Synthetic Metadata Subject"
properties.keywords = "synthetic-metadata-keyword"
properties.comments = "Synthetic Metadata Comment"

document.add_paragraph("Synthetic Title", style="Title")
document.add_paragraph("Synthetic Heading", style="Heading 1")
body = document.add_paragraph("Synthetic body text with ")
ligature_run = body.add_run("ligature fi")
ligatures = OxmlElement("w14:ligatures")
ligatures.set(f"{{{W14}}}val", "standardContextual")
ligature_run._r.get_or_add_rPr().append(ligatures)
body.add_run(".")
document.add_paragraph("First synthetic item", style="List Bullet")
document.add_paragraph("Second synthetic item", style="List Bullet")
east_asian = document.add_paragraph().add_run("合成テキストの段落です。")
fonts = OxmlElement("w:rFonts")
fonts.set(qn("w:hint"), "eastAsia")
east_asian._r.get_or_add_rPr().append(fonts)
table = document.add_table(rows=2, cols=2, style="Light Shading Accent 1")
for row_index, row in enumerate(table.rows):
    for column_index, cell in enumerate(row.cells):
        cell.text = f"Synthetic cell {row_index}{column_index}"
document.save(sys.argv[1])
